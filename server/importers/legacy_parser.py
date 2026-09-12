from __future__ import annotations

import hashlib
import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from excel_profiles import detect_mapping, load_profile, rows_from_preview

DATE_PATTERNS = [
    re.compile(r"\b(\d{1,2})[.\-_](\d{1,2})[.\-_](20\d{2})\b"),
    re.compile(r"\b(20\d{2})[.\-_](\d{1,2})[.\-_](\d{1,2})\b"),
]

NOISE_WORDS = {
    "sozlesme", "sözleşme", "kira", "sevkiyat", "tablosu", "tablo",
    "giden", "gelen", "iade", "fatura", "excel", "xlsx", "pdf",
    "jpg", "jpeg", "png", "doc", "docx",
}

def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def infer_kind(path: Path) -> str:
    name = path.stem.lower()
    ext = path.suffix.lower()

    if ext in {".xlsx", ".xlsm"}:
        return "excel"
    if "sözleş" in name or "sozles" in name:
        return "contract"
    if any(x in name for x in ("sevkiyat", "giden", "gelen", "teslim", "iade")):
        return "delivery_table"
    if "fatura" in name:
        return "invoice"
    return "other"

def infer_date_from_name(path: Path) -> date | None:
    text = path.stem
    for idx, pattern in enumerate(DATE_PATTERNS):
        m = pattern.search(text)
        if not m:
            continue
        try:
            if idx == 0:
                day, month, year = map(int, m.groups())
            else:
                year, month, day = map(int, m.groups())
            return date(year, month, day)
        except ValueError:
            pass
    return None

def infer_customer_from_name(path: Path) -> str | None:
    text = path.stem
    for pattern in DATE_PATTERNS:
        text = pattern.sub(" ", text)

    text = re.sub(r"[_\-]+", " ", text)
    words = []
    for w in text.split():
        clean = w.strip(" .()[]{}").lower()
        if not clean or clean in NOISE_WORDS:
            continue
        words.append(w.strip(" .()[]{}"))

    candidate = " ".join(words).strip()
    return candidate or None

def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)

def inspect_excel(path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    wb = load_workbook(
        path,
        read_only=True,
        data_only=True,
    )

    sheets = []
    date_candidates: list[str] = []
    text_candidates: list[str] = []

    for ws in wb.worksheets[:10]:
        preview = []
        for row_index, row in enumerate(
            ws.iter_rows(min_row=1, max_row=60, values_only=True),
            start=1,
        ):
            values = [_json_value(v) for v in row[:20]]
            if any(v not in (None, "") for v in values):
                preview.append({
                    "row": row_index,
                    "values": values,
                })

            for v in row[:20]:
                if isinstance(v, (datetime, date)):
                    date_candidates.append(v.date().isoformat() if isinstance(v, datetime) else v.isoformat())
                elif isinstance(v, str):
                    t = v.strip()
                    if 3 <= len(t) <= 100:
                        text_candidates.append(t)

            if len(preview) >= 25:
                break

        mapping = detect_mapping(
            rows_from_preview(preview),
            load_profile(),
        )

        sheets.append({
            "name": ws.title,
            "preview": preview,
            "mapping_suggestion": mapping,
        })

    return {
        "sheet_count": len(wb.sheetnames),
        "sheet_names": wb.sheetnames,
        "sheets": sheets,
        "date_candidates": date_candidates[:30],
        "text_candidates": text_candidates[:100],
    }

def parse_file(path: Path) -> dict[str, Any]:
    path = path.resolve()
    result: dict[str, Any] = {
        "file_name": path.name,
        "extension": path.suffix.lower(),
        "size_bytes": path.stat().st_size,
        "sha256": sha256_file(path),
        "source_kind": infer_kind(path),
        "detected_customer_name": infer_customer_from_name(path),
        "detected_date": (
            infer_date_from_name(path).isoformat()
            if infer_date_from_name(path)
            else None
        ),
        "parser_version": 1,
    }

    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        try:
            result["excel"] = inspect_excel(path)
        except Exception as exc:
            result["excel_error"] = str(exc)

    # PDF/fotoğraflarda bu aşamada OCR yapılmaz.
    # Dosya staging'e alınır; içerik eşleştirme daha sonra insan kontrolüyle yapılır.
    return result

def dumps(result: dict[str, Any]) -> str:
    return json.dumps(result, ensure_ascii=False, indent=2)
