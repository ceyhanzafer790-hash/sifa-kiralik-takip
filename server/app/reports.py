from __future__ import annotations

import csv
import io
from datetime import date
from typing import Any

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font

from .db import db
from .security import current_user, require_admin

router = APIRouter(prefix="/reports", tags=["reports"])

UNIT_LABELS = {
    "piece": "Adet",
    "sheet": "Levha",
    "meter": "Metre",
    "square_meter": "m²",
    "cubic_meter": "m³",
    "kilogram": "kg",
    "liter": "Litre",
    "set": "Takım",
}

RENTAL_HEADERS = [
    ("customer_name", "Müşteri"),
    ("address_label", "Şantiye"),
    ("original_outbound_date", "İlk Çıkış"),
    ("product_name", "Malzeme"),
    ("unit_label", "Birim"),
    ("initial_quantity", "Gönderilen"),
    ("returned_quantity", "İade"),
    ("remaining_quantity", "Kirada Kalan"),
    ("current_rate", "Güncel Kira Fiyatı"),
    ("current_rate_type_label", "Fiyat Tipi"),
    ("invoice_preference_label", "Fatura"),
    ("rental_status_label", "Durum"),
]

STOCK_HEADERS = [
    ("product_name", "Malzeme"),
    ("unit_label", "Birim"),
    ("available_quantity", "Kullanılabilir"),
    ("repair_quantity", "Tamirlik"),
    ("scrap_quantity", "Hurda"),
    ("lost_quantity", "Kayıp"),
    ("stock_confidence_label", "Stok Güveni"),
    ("last_count_at", "Son Fiziksel Sayım"),
]

SALES_HEADERS = [
    ("customer_name", "Müşteri"),
    ("sale_date", "Satış Tarihi"),
    ("product_name", "Malzeme"),
    ("unit_label", "Birim"),
    ("quantity", "Miktar"),
    ("unit_price", "Birim Fiyat"),
    ("line_total", "Toplam"),
    ("note", "Not"),
]


AUDIT_HEADERS = [
    ("created_at", "Tarih/Saat"),
    ("user_name", "Kullanıcı"),
    ("user_email", "E-posta"),
    ("entity_type", "Kayıt Türü"),
    ("entity_id", "Kayıt ID"),
    ("action", "İşlem"),
    ("payload_text", "Detay"),
]

BILLING_HEADERS = [
    ("customer_name", "Müşteri"),
    ("address_label", "Şantiye"),
    ("renewal_date", "Kira Dönemi"),
    ("invoice_status_label", "Fatura Durumu"),
    ("invoice_date", "Fatura Tarihi"),
    ("invoice_no", "Fatura No"),
    ("payment_due_date", "Ödeme Vadesi"),
    ("billed_amount", "Fatura Tutarı"),
    ("payment_status_label", "Tahsilat Durumu"),
    ("paid_amount", "Tahsil Edilen"),
    ("balance_amount", "Kalan"),
]


def _rentals(
    status: str = "active",
    customer_id: str | None = None,
    from_date: date | None = None,
    to_date: date | None = None,
) -> list[dict[str, Any]]:
    if status not in ("active", "closed", "all"):
        raise HTTPException(
            status_code=400,
            detail="Geçersiz kiralama durumu.",
        )

    conditions = []
    params: list[Any] = []

    if status != "all":
        conditions.append("rental_status = %s")
        params.append(status)
    if customer_id:
        conditions.append("customer_id = %s")
        params.append(customer_id)
    if from_date:
        conditions.append("original_outbound_date >= %s")
        params.append(from_date)
    if to_date:
        conditions.append("original_outbound_date <= %s")
        params.append(to_date)

    where = "where " + " and ".join(conditions) if conditions else ""

    with db() as (_, cur):
        cur.execute(
            f"""
            select *
            from v_rental_item_status
            {where}
            order by customer_name, original_outbound_date, product_name
            """,
            tuple(params),
        )
        rows = cur.fetchall()

    result = []
    for row in rows:
        d = dict(row)
        d["unit_label"] = UNIT_LABELS.get(d["unit"], d["unit"])
        d["current_rate_type_label"] = {
            "per_unit_monthly": "Birim başına aylık",
            "fixed_monthly": "Sabit aylık",
            None: "",
        }.get(d["current_rate_type"], d["current_rate_type"] or "")
        d["invoice_preference_label"] = (
            "Faturalı"
            if d["invoice_preference"] == "invoice_required"
            else "Faturasız"
        )
        d["rental_status_label"] = (
            "Aktif" if d["rental_status"] == "active" else "Kapalı"
        )
        result.append(d)
    return result


def _stock(
    confidence: str | None = None,
    category: str | None = None,
) -> list[dict[str, Any]]:
    if confidence and confidence not in (
        "counted",
        "estimated",
        "unknown",
    ):
        raise HTTPException(
            status_code=400,
            detail="Geçersiz stok güven seviyesi.",
        )

    conditions = []
    params: list[Any] = []

    if confidence:
        conditions.append("stock_confidence = %s")
        params.append(confidence)
    if category:
        conditions.append("category = %s")
        params.append(category)

    where = "where " + " and ".join(conditions) if conditions else ""

    with db() as (_, cur):
        cur.execute(
            f"""
            select *
            from v_stock_summary
            {where}
            order by product_name
            """,
            tuple(params),
        )
        rows = cur.fetchall()

    result = []
    for row in rows:
        d = dict(row)
        d["unit_label"] = UNIT_LABELS.get(d["unit"], d["unit"])
        d["stock_confidence_label"] = {
            "counted": "Fiziksel sayım",
            "estimated": "Tahmini / hareketlerden",
            "unknown": "Bilinmiyor",
        }.get(d["stock_confidence"], d["stock_confidence"])
        result.append(d)
    return result


def _sales(
    customer_id: str | None = None,
    from_date: date | None = None,
    to_date: date | None = None,
) -> list[dict[str, Any]]:
    conditions = ["s.status = 'completed'"]
    params: list[Any] = []

    if customer_id:
        conditions.append("s.customer_id = %s")
        params.append(customer_id)
    if from_date:
        conditions.append("s.sale_date >= %s")
        params.append(from_date)
    if to_date:
        conditions.append("s.sale_date <= %s")
        params.append(to_date)

    where = "where " + " and ".join(conditions)

    with db() as (_, cur):
        cur.execute(
            f"""
            select
              c.name as customer_name,
              s.sale_date,
              p.name as product_name,
              p.unit,
              si.quantity,
              si.unit_price,
              si.line_total,
              s.note
            from sales s
            join customers c on c.id = s.customer_id
            join sale_items si on si.sale_id = s.id
            join products p on p.id = si.product_id
            {where}
            order by s.sale_date desc, c.name, p.name
            """,
            tuple(params),
        )
        rows = cur.fetchall()

    result = []
    for row in rows:
        d = dict(row)
        d["unit_label"] = UNIT_LABELS.get(d["unit"], d["unit"])
        result.append(d)
    return result


def _billing(
    from_date: date | None = None,
    to_date: date | None = None,
    customer_id: str | None = None,
    invoice_status: str | None = None,
    payment_status: str | None = None,
) -> list[dict[str, Any]]:
    if invoice_status and invoice_status not in (
        "not_required",
        "pending",
        "issued",
    ):
        raise HTTPException(
            status_code=400,
            detail="Geçersiz fatura durumu.",
        )

    if payment_status and payment_status not in (
        "pending",
        "partial",
        "paid",
    ):
        raise HTTPException(
            status_code=400,
            detail="Geçersiz tahsilat durumu.",
        )

    conditions = []
    params: list[Any] = []

    if from_date:
        conditions.append("renewal_date >= %s")
        params.append(from_date)
    if to_date:
        conditions.append("renewal_date <= %s")
        params.append(to_date)
    if customer_id:
        conditions.append("customer_id = %s")
        params.append(customer_id)
    if invoice_status:
        conditions.append("invoice_status = %s")
        params.append(invoice_status)
    if payment_status:
        conditions.append("payment_status = %s")
        params.append(payment_status)

    where = "where " + " and ".join(conditions) if conditions else ""

    with db() as (_, cur):
        cur.execute(
            f"""
            select *
            from v_billing_report
            {where}
            order by renewal_date desc, customer_name
            """,
            tuple(params),
        )
        rows = cur.fetchall()

    result = []
    for row in rows:
        d = dict(row)
        d["invoice_status_label"] = {
            "not_required": "Fatura yok",
            "pending": "Fatura bekliyor",
            "issued": "Fatura kesildi",
        }.get(d["invoice_status"], d["invoice_status"])
        d["payment_status_label"] = {
            "pending": "Bekliyor",
            "partial": "Kısmi",
            "paid": "Ödendi",
        }.get(d["payment_status"], d["payment_status"])
        result.append(d)
    return result


def _cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, date):
        return value.isoformat()
    return value


def _csv_bytes(
    rows: list[dict[str, Any]],
    headers: list[tuple[str, str]],
) -> bytes:
    stream = io.StringIO()
    writer = csv.writer(stream, delimiter=";")
    writer.writerow([label for _, label in headers])

    for row in rows:
        writer.writerow(
            [_cell(row.get(key)) for key, _ in headers]
        )

    return ("\ufeff" + stream.getvalue()).encode("utf-8")


def _style_sheet(ws) -> None:
    if ws.max_row >= 1:
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    for column_cells in ws.columns:
        max_len = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        ws.column_dimensions[column_cells[0].column_letter].width = min(
            max(max_len + 2, 10),
            42,
        )


def _xlsx_bytes(
    title: str,
    rows: list[dict[str, Any]],
    headers: list[tuple[str, str]],
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append([label for _, label in headers])

    for row in rows:
        ws.append([_cell(row.get(key)) for key, _ in headers])

    _style_sheet(ws)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def _response(
    report_name: str,
    title: str,
    rows: list[dict[str, Any]],
    headers: list[tuple[str, str]],
    fmt: str,
) -> Response:
    stamp = date.today().isoformat()
    filename = f"sifa_{report_name}_{stamp}.{fmt}"

    if fmt == "csv":
        content = _csv_bytes(rows, headers)
        media = "text/csv; charset=utf-8"
    elif fmt == "xlsx":
        content = _xlsx_bytes(title, rows, headers)
        media = (
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    else:
        raise HTTPException(
            status_code=400,
            detail="Format csv veya xlsx olmalı.",
        )

    return Response(
        content=content,
        media_type=media,
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "X-Report-Rows": str(len(rows)),
        },
    )


@router.get("/rentals")
def rentals_report(
    format: str = Query(default="xlsx", pattern="^(xlsx|csv)$"),
    status: str = Query(default="active"),
    customer_id: str | None = None,
    from_date: date | None = None,
    to_date: date | None = None,
    user=Depends(current_user),
):
    _validate_range(from_date, to_date)
    rows = _rentals(
        status=status,
        customer_id=customer_id,
        from_date=from_date,
        to_date=to_date,
    )
    return _response(
        "rentals",
        "Kiralama Takibi",
        rows,
        RENTAL_HEADERS,
        format,
    )


@router.get("/stock")
def stock_report(
    format: str = Query(default="xlsx", pattern="^(xlsx|csv)$"),
    confidence: str | None = None,
    category: str | None = None,
    user=Depends(current_user),
):
    return _response(
        "stock",
        "Stok Özeti",
        _stock(confidence=confidence, category=category),
        STOCK_HEADERS,
        format,
    )


@router.get("/sales")
def sales_report(
    format: str = Query(default="xlsx", pattern="^(xlsx|csv)$"),
    customer_id: str | None = None,
    from_date: date | None = None,
    to_date: date | None = None,
    user=Depends(current_user),
):
    _validate_range(from_date, to_date)

    return _response(
        "sales",
        "Satışlar",
        _sales(
            customer_id=customer_id,
            from_date=from_date,
            to_date=to_date,
        ),
        SALES_HEADERS,
        format,
    )


@router.get("/billing")
def billing_report(
    format: str = Query(default="xlsx", pattern="^(xlsx|csv)$"),
    from_date: date | None = None,
    to_date: date | None = None,
    customer_id: str | None = None,
    invoice_status: str | None = None,
    payment_status: str | None = None,
    user=Depends(require_admin),
):
    _validate_range(from_date, to_date)

    return _response(
        "billing",
        "Fatura ve Tahsilat",
        _billing(
            from_date=from_date,
            to_date=to_date,
            customer_id=customer_id,
            invoice_status=invoice_status,
            payment_status=payment_status,
        ),
        BILLING_HEADERS,
        format,
    )


@router.get("/customer-statement")
def customer_statement(
    customer_id: str,
    from_date: date | None = None,
    to_date: date | None = None,
    user=Depends(require_admin),
):
    _validate_range(from_date, to_date)

    with db() as (_, cur):
        cur.execute(
            """
            select id, name, phone, notes
            from customers
            where id = %s and deleted_at is null
            """,
            (customer_id,),
        )
        customer = cur.fetchone()

    if not customer:
        raise HTTPException(
            status_code=404,
            detail="Müşteri bulunamadı.",
        )

    rentals = _rentals(
        status="all",
        customer_id=customer_id,
        from_date=from_date,
        to_date=to_date,
    )
    billing = _billing(
        customer_id=customer_id,
        from_date=from_date,
        to_date=to_date,
    )

    sale_conditions = [
        "s.customer_id = %s",
        "s.status = 'completed'",
    ]
    sale_params: list[Any] = [customer_id]

    if from_date:
        sale_conditions.append("s.sale_date >= %s")
        sale_params.append(from_date)
    if to_date:
        sale_conditions.append("s.sale_date <= %s")
        sale_params.append(to_date)

    with db() as (_, cur):
        cur.execute(
            f"""
            select
              s.id,
              s.sale_date,
              s.note,
              p.name as product_name,
              p.unit,
              si.quantity,
              si.unit_price,
              si.line_total
            from sales s
            join sale_items si on si.sale_id = s.id
            join products p on p.id = si.product_id
            where {" and ".join(sale_conditions)}
            order by s.sale_date desc, p.name
            """,
            tuple(sale_params),
        )
        sales = cur.fetchall()

    total_billed = sum(float(r["billed_amount"] or 0) for r in billing)
    total_paid = sum(float(r["paid_amount"] or 0) for r in billing)
    balance = total_billed - total_paid
    sales_total = sum(float(r["line_total"] or 0) for r in sales)

    wb = Workbook()
    summary = wb.active
    summary.title = "Özet"
    summary.append(["Alan", "Değer"])
    summary.append(["Müşteri", customer["name"]])
    summary.append(["Telefon", customer["phone"] or ""])
    summary.append([
        "Dönem",
        (
            f"{from_date or 'Başlangıç'} - {to_date or 'Bugün'}"
            if from_date or to_date
            else "Tüm dönem"
        ),
    ])
    summary.append(["Kiralık kalem", len(rentals)])
    summary.append(["Kira Faturalandırılan", total_billed])
    summary.append(["Kira Tahsil Edilen", total_paid])
    summary.append(["Kira Kalan Bakiye", balance])
    summary.append(["Satış Toplamı", sales_total])
    summary.append([
        "Not",
        "Satış toplamı kira alacak bakiyesine dahil edilmemiştir.",
    ])
    _style_sheet(summary)

    ws_rental = wb.create_sheet("Kiralamalar")
    ws_rental.append([label for _, label in RENTAL_HEADERS])
    for row in rentals:
        ws_rental.append(
            [_cell(row.get(key)) for key, _ in RENTAL_HEADERS]
        )
    _style_sheet(ws_rental)

    ws_billing = wb.create_sheet("Fatura-Tahsilat")
    ws_billing.append([label for _, label in BILLING_HEADERS])
    for row in billing:
        ws_billing.append(
            [_cell(row.get(key)) for key, _ in BILLING_HEADERS]
        )
    _style_sheet(ws_billing)

    ws_sales = wb.create_sheet("Satışlar")
    sale_headers = [
        ("sale_date", "Satış Tarihi"),
        ("product_name", "Malzeme"),
        ("unit_label", "Birim"),
        ("quantity", "Miktar"),
        ("unit_price", "Birim Fiyat"),
        ("line_total", "Toplam"),
        ("note", "Not"),
    ]
    ws_sales.append([label for _, label in sale_headers])

    for raw in sales:
        row = dict(raw)
        row["unit_label"] = UNIT_LABELS.get(row["unit"], row["unit"])
        ws_sales.append(
            [_cell(row.get(key)) for key, _ in sale_headers]
        )
    _style_sheet(ws_sales)

    output = io.BytesIO()
    wb.save(output)

    safe_name = "".join(
        c if c.isalnum() else "_"
        for c in customer["name"]
    ).strip("_") or "musteri"

    return Response(
        content=output.getvalue(),
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; filename="sifa_ekstre_{safe_name}.xlsx"'
            ),
        },
    )


def _validate_range(
    from_date: date | None,
    to_date: date | None,
) -> None:
    if from_date and to_date and from_date > to_date:
        raise HTTPException(
            status_code=400,
            detail="Başlangıç tarihi bitiş tarihinden sonra olamaz.",
        )

@router.get("/audit")
def audit_report(
    format: str = Query(default="xlsx", pattern="^(xlsx|csv)$"),
    user_id: str | None = None,
    entity_type: str | None = None,
    action: str | None = None,
    from_at: str | None = None,
    to_at: str | None = None,
    q: str | None = Query(default=None, max_length=120),
    user=Depends(require_admin),
):
    conditions = []
    params: list[Any] = []

    if user_id:
        conditions.append("a.user_id = %s")
        params.append(user_id)

    if entity_type:
        conditions.append("a.entity_type = %s")
        params.append(entity_type)

    if action:
        conditions.append("a.action = %s")
        params.append(action)

    if from_at:
        conditions.append("a.created_at >= %s::timestamptz")
        params.append(from_at)

    if to_at:
        conditions.append("a.created_at <= %s::timestamptz")
        params.append(to_at)

    if q and q.strip():
        value = f"%{q.strip()}%"
        conditions.append(
            """
            (
              coalesce(u.full_name, '') ilike %s
              or coalesce(u.email, '') ilike %s
              or a.entity_type ilike %s
              or a.entity_id ilike %s
              or a.action ilike %s
              or coalesce(a.payload::text, '') ilike %s
            )
            """
        )
        params.extend([value] * 6)

    where = "where " + " and ".join(conditions) if conditions else ""

    with db() as (_, cur):
        cur.execute(
            f"""
            select
              a.created_at,
              coalesce(u.full_name, 'Sistem') as user_name,
              coalesce(u.email, '') as user_email,
              a.entity_type,
              a.entity_id,
              a.action,
              coalesce(a.payload::text, '') as payload_text
            from audit_logs a
            left join app_users u on u.id = a.user_id
            {where}
            order by a.created_at desc, a.id desc
            """,
            tuple(params),
        )
        rows = [dict(r) for r in cur.fetchall()]

    return _response(
        "audit",
        "İşlem Geçmişi",
        rows,
        AUDIT_HEADERS,
        format,
    )
